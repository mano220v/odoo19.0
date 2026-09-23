import base64
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError
try:
    from pdf2docx import Converter as Pdf2DocxConverter
except ImportError:
    Pdf2DocxConverter = None
try:
    from pypdf import PdfReader
except ImportError:
    from PyPDF2 import PdfReader
from docx import Document


class DocumentConversion(models.Model):
    _name = 'ow.document.conversion'
    _description = 'PDF and Word Conversion'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'create_date desc, id desc'
    name = fields.Char(required=True, default=lambda self: _('New Conversion'), tracking=True)
    user_id = fields.Many2one('res.users', required=True, default=lambda self: self.env.user, index=True)
    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company, index=True)
    input_file = fields.Binary(required=True, attachment=True)
    input_filename = fields.Char(required=True)
    input_size = fields.Integer(compute='_compute_input_info', store=True)
    direction = fields.Selection([('pdf_docx', 'PDF → DOCX'), ('word_pdf', 'DOC/DOCX → PDF')], compute='_compute_input_info', store=True)
    state = fields.Selection([('draft', 'Ready'), ('converting', 'Converting'), ('done', 'Completed'), ('failed', 'Failed')], default='draft', required=True, tracking=True, index=True)
    output_file = fields.Binary(attachment=True, readonly=True)
    output_filename = fields.Char(readonly=True)
    output_size = fields.Integer(readonly=True)
    engine = fields.Selection([('pdf2docx', 'Enhanced PDF Layout Engine'), ('text', 'Free Text Extraction'), ('libreoffice', 'LibreOffice')], readonly=True)
    duration = fields.Float(readonly=True, digits=(12, 3))
    page_count = fields.Integer(readonly=True)
    error_message = fields.Text(readonly=True)
    converted_at = fields.Datetime(readonly=True)
    keep_until = fields.Date(default=lambda self: fields.Date.add(fields.Date.today(), days=30), help='Output may be removed by automatic cleanup after this date.')
    notes = fields.Text()

    @api.depends('input_file', 'input_filename')
    def _compute_input_info(self):
        for record in self:
            raw = base64.b64decode(record.input_file or b'')
            record.input_size = len(raw)
            suffix = Path(record.input_filename or '').suffix.lower()
            record.direction = 'pdf_docx' if suffix == '.pdf' else 'word_pdf' if suffix in ('.doc', '.docx') else False

    @api.constrains('input_file', 'input_filename')
    def _check_input(self):
        for record in self:
            suffix = Path(record.input_filename or '').suffix.lower()
            if suffix not in ('.pdf', '.doc', '.docx'):
                raise ValidationError(_('Upload a PDF, DOC, or DOCX file.'))
            if record.input_size > 25 * 1024 * 1024:
                raise ValidationError(_('The maximum file size is 25 MB.'))
            raw = base64.b64decode(record.input_file or b'')
            if suffix == '.pdf' and (not raw.startswith(b'%PDF')):
                raise ValidationError(_('The uploaded file is not a valid PDF.'))
            if suffix == '.docx' and (not raw.startswith(b'PK')):
                raise ValidationError(_('The uploaded file is not a valid DOCX package.'))

    def action_convert(self):
        self.ensure_one()
        self._check_input()
        self.write({'state': 'converting', 'error_message': False, 'output_file': False, 'output_filename': False})
        started = time.monotonic()
        try:
            raw = base64.b64decode(self.input_file)
            with tempfile.TemporaryDirectory(prefix='odoo_doc_convert_') as folder:
                safe_stem = ''.join((ch for ch in Path(self.input_filename).stem if ch.isalnum() or ch in '-_ ')).strip()[:80] or 'document'
                suffix = Path(self.input_filename).suffix.lower()
                source = Path(folder) / (safe_stem + suffix)
                source.write_bytes(raw)
                if self.direction == 'pdf_docx':
                    (output, engine, pages) = self._pdf_to_docx(source, Path(folder), safe_stem)
                    filename = safe_stem + '.docx'
                else:
                    (output, engine, pages) = self._word_to_pdf(source, Path(folder), safe_stem)
                    filename = safe_stem + '.pdf'
                result = output.read_bytes()
            self.write({'state': 'done', 'output_file': base64.b64encode(result), 'output_filename': filename, 'output_size': len(result), 'engine': engine, 'page_count': pages, 'duration': time.monotonic() - started, 'converted_at': fields.Datetime.now()})
        except Exception as exc:
            self.write({'state': 'failed', 'duration': time.monotonic() - started, 'error_message': str(exc)[:4000]})
            raise UserError(_('Conversion failed: %s') % exc) from exc
        return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {'title': _('Conversion completed'), 'message': _('%s is ready to download.') % self.output_filename, 'type': 'success', 'sticky': False, 'next': {'type': 'ir.actions.client', 'tag': 'reload'}}}

    def _pdf_to_docx(self, source, folder, stem):
        output = folder / (stem + '.docx')
        reader = PdfReader(str(source))
        pages = len(reader.pages)
        if Pdf2DocxConverter:
            converter = Pdf2DocxConverter(str(source))
            try:
                converter.convert(str(output))
            finally:
                converter.close()
            return (output, 'pdf2docx', pages)
        document = Document()
        document.core_properties.title = stem
        document.add_heading(stem, 0)
        for (index, page) in enumerate(reader.pages):
            if index:
                document.add_page_break()
            document.add_heading(_('Page %s') % (index + 1), level=2)
            text = page.extract_text() or ''
            for paragraph in text.splitlines():
                if paragraph.strip():
                    document.add_paragraph(paragraph.strip())
        document.save(str(output))
        return (output, 'text', pages)

    def _word_to_pdf(self, source, folder, stem):
        executable = shutil.which('libreoffice') or shutil.which('soffice')
        if not executable:
            raise UserError(_('LibreOffice is required for Word-to-PDF conversion. Install it on the Odoo server.'))
        profile = folder / 'lo_profile'
        command = [executable, '--headless', '--nologo', '--nodefault', '--nolockcheck', '-env:UserInstallation=file://%s' % profile, '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(folder), str(source)]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=120, check=False)
        output = folder / (source.stem + '.pdf')
        if completed.returncode or not output.exists():
            raise UserError(_('LibreOffice could not convert this document: %s') % (completed.stderr or completed.stdout or _('unknown error')))
        pages = len(PdfReader(str(output)).pages)
        return (output, 'libreoffice', pages)

    def action_reset(self):
        self.write({'state': 'draft', 'output_file': False, 'output_filename': False, 'output_size': 0, 'engine': False, 'duration': 0, 'page_count': 0, 'error_message': False, 'converted_at': False})
        return True

    def action_new_conversion(self):
        return {'type': 'ir.actions.act_window', 'name': _('New Conversion'), 'res_model': self._name, 'views': [(False, 'form')], 'target': 'current', 'context': {'default_name': _('New Conversion')}}

    @api.model
    def _cron_cleanup_outputs(self):
        expired = self.search([('keep_until', '<', fields.Date.today()), ('output_file', '!=', False)])
        expired.write({'output_file': False, 'output_filename': False, 'output_size': 0})

    @api.model
    def get_engine_status(self):
        return {'libreoffice': bool(shutil.which('libreoffice') or shutil.which('soffice')), 'pdf2docx': Pdf2DocxConverter is not None, 'fallback': True}
